from docx import Document

from kpis import add_handle_time, add_locat_time, load_chats, overall_kpis, kpis_by


def add_kpi_table(doc, table_df, heading):
    doc.add_heading(heading, level=1)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"

    header = table.rows[0].cells
    header[0].text = table_df.index.name.replace("_", "").capitalize()
    header[1].text = "Chats"
    header[2].text = "AHT (min)"

    for name, row in table_df.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(name)
        cells[1].text = str(int(row["volume"]))
        cells[2].text = f"{row['aht_minutes']:.1f}"


def build_report(path="kpi_report.docx"):
    df = add_locat_time(add_handle_time(load_chats()))
    summary = overall_kpis(df)

    doc = Document()
    doc.add_heading("Support KPI Report", level=0)
    doc.add_paragraph(
        f"Period:  {df['date'].min()}   to   {df['date'].max()}(Tbilisi time)")

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(f"Total chats: {summary['volume']}")
    doc.add_paragraph(f"Average handle time: {summary['aht_minutes']} min")
    doc.add_paragraph(f"Median handle time: {summary['median_minutes']} min")

    add_kpi_table(doc, kpis_by(df, "agent_name"), "by agent")
    add_kpi_table(doc, kpis_by(df, "channel"), "by channel")
    add_kpi_table(doc, kpis_by(df, "date").sort_index(), "by day")

    doc.save(path)

    print("Report saved:", path)


if __name__ == "__main__":
    build_report()
